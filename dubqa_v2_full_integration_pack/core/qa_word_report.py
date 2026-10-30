"""Build a Word report exclusively from the QA_Results.xlsx workbook."""
from __future__ import annotations
from pathlib import Path
from openpyxl import load_workbook
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def _rows(ws):
    values = list(ws.iter_rows(values_only=True))
    if not values: return [], []
    headers = [str(x) if x is not None else "" for x in values[0]]
    return headers, [dict(zip(headers, row)) for row in values[1:]]

def _shade(cell, color):
    pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), color); pr.append(shd)

def _add_table(doc, headers, rows, max_rows=30):
    if not headers: return
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        t.rows[0].cells[i].text = str(h)
        _shade(t.rows[0].cells[i], "17365D")
        for r in t.rows[0].cells[i].paragraphs[0].runs:
            r.font.bold = True; r.font.color.rgb = RGBColor(255,255,255); r.font.size = Pt(8)
    for row in rows[:max_rows]:
        cells = t.add_row().cells
        for i,h in enumerate(headers):
            v = row.get(h, "")
            cells[i].text = "" if v is None else str(v)
            for p in cells[i].paragraphs:
                for r in p.runs: r.font.size = Pt(7.5)
    if len(rows) > max_rows:
        doc.add_paragraph(f"Showing first {max_rows} of {len(rows)} rows. See Excel workbook for full detail.")

def build_from_workbook(xlsx_path: str | Path, output_path: str | Path | None = None) -> Path:
    xlsx_path = Path(xlsx_path)
    if not xlsx_path.exists(): raise FileNotFoundError(xlsx_path)
    wb = load_workbook(xlsx_path, data_only=True, read_only=True)
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(.65); sec.bottom_margin = Inches(.65)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("DUBQA V2 – QUALITY ASSURANCE REPORT")
    run.bold = True; run.font.size = Pt(21); run.font.color.rgb = RGBColor(23,54,93)
    doc.add_paragraph("Generated from the run's QA_Results.xlsx workbook. Excel is the source of truth.")
    meta = []
    if "Run Metadata" in wb.sheetnames:
        meta = list(wb["Run Metadata"].iter_rows(values_only=True))
    if meta:
        doc.add_heading("Run Information", level=1)
        _add_table(doc, ["Field","Value"], [{"Field":r[0],"Value":r[1]} for r in meta if len(r)>1])
    if "Executive Summary" in wb.sheetnames:
        doc.add_heading("Executive Summary", level=1)
        h, rows = _rows(wb["Executive Summary"]); _add_table(doc,h,rows,50)
    sections = [
        ("Vendor Metrics", "Vendor and Metric Results", 25),
        ("Scenario Results", "Scenario Test Results", 35),
        ("Defects", "Defect Register", 35),
        ("Audio Quality", "Audio Quality", 20),
        ("Captions", "Caption QA", 20),
        ("Competitor Benchmark", "Competitor Benchmark", 20),
    ]
    for sheet, heading, limit in sections:
        if sheet in wb.sheetnames:
            h, rows = _rows(wb[sheet])
            doc.add_heading(heading, level=1)
            _add_table(doc,h,rows,limit)
    if "Word Timeline" in wb.sheetnames:
        h, rows = _rows(wb["Word Timeline"])
        doc.add_heading("Word-level Lip-sync Evidence", level=1)
        _add_table(doc,h,rows,40)
    doc.add_heading("Interpretation and Limitations", level=1)
    doc.add_paragraph(
        "PASS means the configured and calibrated checks passed. FAIL means a defect was "
        "detected. REVIEW means human verification or stronger evidence is required. "
        "NOT_MEASURED means required data or a validated measurement method was unavailable. "
        "A mouth-aperture signal is a visual speech-activity proxy and is not phoneme-to-viseme "
        "validation unless the run explicitly records a validated phoneme aligner and viseme map."
    )
    out = Path(output_path) if output_path else xlsx_path.with_name("QA_Report.docx")
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    return out
