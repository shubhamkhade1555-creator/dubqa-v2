import json
from pathlib import Path
from openpyxl import load_workbook
from core.qa_report_bundle import export_run_reports
from core.qa_public_summary import build_sanitized_summary
import pytest

def make_run(tmp_path):
    run = tmp_path / "RUN_001"
    (run / "report").mkdir(parents=True)
    (run / "words").mkdir()
    (run / "measure").mkdir()
    results = {
        "run_id":"RUN_001", "generated":"2026-10-30T10:00:00Z", "mode":"precomputed",
        "source":{"file":"source.mp4"}, "outputs":{
            "dub_01":{"vendor":"TestVendor","values":{"translation_fidelity":90,"onset_drift_ms":40},
                      "extras":{"missing_words":["example"]},"score":{"overall":85,"coverage_pct":70}}
        },
        "scenario_results":[{"scenario_id":"WC-01","category":"completeness","status":"FAIL",
          "severity":"HIGH","expected":"one word","observed":"missing","evidence_start":2.1,"evidence_end":2.4}],
        "defects":[{"defect_id":"D-1","category":"completeness","severity":"HIGH",
          "description":"Missing word","expected":"example","actual":"missing","start":2.1,"end":2.4,
          "evidence":["frame_052.jpg"],"status":"FAIL"}]
    }
    (run / "results.json").write_text(json.dumps(results), encoding="utf-8")
    (run / "words" / "dub_01_words.json").write_text(json.dumps({"words":[{"word":"hello","start":1.0,"end":1.4}]}))
    (run / "measure" / "dub_01.json").write_text(json.dumps({"drifts":[{"onset_ms":40,"tail_ms":-20}]}))
    return run

def test_xlsx_and_docx_generated_from_run(tmp_path):
    run = make_run(tmp_path)
    paths = export_run_reports(run)
    assert Path(paths["xlsx"]).exists()
    assert Path(paths["docx"]).exists()
    wb = load_workbook(paths["xlsx"], data_only=True)
    assert "Word Timeline" in wb.sheetnames
    assert "Scenario Results" in wb.sheetnames
    assert "Defects" in wb.sheetnames
    assert wb["Word Timeline"].max_row == 2

def test_public_summary_requires_approval():
    with pytest.raises(PermissionError):
        build_sanitized_summary({"run_id":"x"}, approved=False)

def test_public_summary_does_not_include_transcript_or_media():
    result = {"run_id":"RUN_12345678","generated":"2026-10-30","source_transcript":"RAW_TRANSCRIPT_SENTINEL",
              "outputs":{"a":{"vendor":"V","values":{"onset_drift_ms":20},"file":"private.mp4"}}}
    data = build_sanitized_summary(result, approved=True)
    serialized = json.dumps(data)
    assert "RAW_TRANSCRIPT_SENTINEL" not in serialized
    assert "private.mp4" not in serialized
