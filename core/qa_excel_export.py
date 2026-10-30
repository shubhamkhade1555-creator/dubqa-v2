"""Export one DubQA run to an evidence-oriented XLSX workbook.

This module is intentionally tolerant of optional fields from older results.json files.
Missing evidence is exported as NOT_MEASURED; it is never converted into PASS.
"""
from __future__ import annotations
import json
from pathlib import Path
from datetime import datetime, timezone
from typing import Any
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

STATUS = {"PASS", "FAIL", "REVIEW", "NOT_MEASURED", "NOT_TESTED"}

def _read_json(path: Path, default):
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError, TypeError):
        return default

def _status(value):
    return value if value in STATUS else "NOT_MEASURED"

def _flatten_metrics(results):
    rows = []
    for oid, out in (results.get("outputs") or {}).items():
        vals = out.get("values") or {}
        extras = out.get("extras") or {}
        score = out.get("score") or {}
        rows.append({
            "Output ID": oid, "Vendor": out.get("vendor", ""),
            "Overall Score": score.get("overall"), "Coverage %": score.get("coverage_pct"),
            "Rank": out.get("overall_rank"),
            "Translation Fidelity %": vals.get("translation_fidelity"),
            "Word Count Diff %": vals.get("word_count_diff"),
            "Timing Drift ms": vals.get("timing_drift_ms"),
            "Mouth Onset Drift ms": vals.get("onset_drift_ms"),
            "Mouth Tail Drift ms": vals.get("tail_drift_ms"),
            "Word Lip-sync Pass %": vals.get("word_pass_rate_pct"),
            "False Movement %": vals.get("false_movement_pct"),
            "Missing Sentences": vals.get("missing_sentences"),
            "Missing Words": ", ".join(map(str, extras.get("missing_words") or [])),
            "Extra Words": ", ".join(map(str, extras.get("extra_words") or [])),
            "Duration s": extras.get("duration_s"),
            "A/V Offset ms": extras.get("av_offset_ms"),
            "Status": _status(out.get("status")),
        })
    return rows

def _word_rows(run_dir: Path, results):
    rows = []
    source = _read_json(run_dir / "words" / "source_words.json",
                        _read_json(run_dir / "precomputed" / "source_words.json", {}))
    src_words = source.get("words", []) if isinstance(source, dict) else []
    for oid, out in (results.get("outputs") or {}).items():
        data = _read_json(run_dir / "words" / f"{oid}_words.json",
                          _read_json(run_dir / "precomputed" / f"{oid}_words.json", {}))
        words = data.get("words", []) if isinstance(data, dict) else []
        measure = _read_json(run_dir / "measure" / f"{oid}.json", {})
        drifts = measure.get("drifts", []) or []
        by_index = {i: d for i, d in enumerate(drifts)}
        for i, w in enumerate(words):
            d = by_index.get(i, {})
            rows.append({
                "Output ID": oid, "Vendor": out.get("vendor", ""),
                "Word Index": i + 1, "Word": w.get("word"),
                "Word Start s": w.get("start"), "Word End s": w.get("end"),
                "Speaker ID": w.get("speaker_id"), "Face Track ID": w.get("face_track_id"),
                "Phoneme Data": json.dumps(w.get("phonemes"), ensure_ascii=False) if w.get("phonemes") else "NOT_MEASURED",
                "Mouth Start s": w.get("mouth_start"), "Mouth Peak s": w.get("mouth_peak"),
                "Mouth End s": w.get("mouth_end"),
                "Onset Offset ms": d.get("onset_ms", w.get("start_offset_ms")),
                "Tail Offset ms": d.get("tail_ms", w.get("end_offset_ms")),
                "Tracking Confidence": w.get("mouth_tracking_confidence"),
                "Alignment Confidence": w.get("phoneme_alignment_confidence"),
                "Lip-sync Status": _status(w.get("lip_sync_status")),
                "Evidence": w.get("evidence_ref", ""),
            })
    return rows

def _scenario_rows(results):
    rows = []
    for item in results.get("scenario_results", []) or []:
        rows.append({
            "Scenario ID": item.get("scenario_id"), "Category": item.get("category"),
            "Title": item.get("title"), "Status": _status(item.get("status")),
            "Severity": item.get("severity"), "Expected": item.get("expected"),
            "Observed": item.get("observed"), "Evidence Start s": item.get("evidence_start"),
            "Evidence End s": item.get("evidence_end"), "Confidence": item.get("confidence"),
            "Evidence Refs": ", ".join(map(str, item.get("evidence_refs") or [])),
            "Notes": item.get("notes"),
        })
    return rows

def _defect_rows(results):
    rows = []
    for d in results.get("defects", []) or []:
        rows.append({
            "Defect ID": d.get("defect_id"), "Category": d.get("category"),
            "Severity": d.get("severity"), "Description": d.get("description"),
            "Expected": d.get("expected"), "Actual": d.get("actual"),
            "Start s": d.get("start"), "End s": d.get("end"),
            "Evidence": ", ".join(map(str, d.get("evidence") or [])),
            "Status": _status(d.get("status")), "Reviewer Note": d.get("reviewer_note"),
        })
    # Preserve omissions already surfaced by the legacy pipeline.
    if not rows:
        for oid, out in (results.get("outputs") or {}).items():
            extras = out.get("extras") or {}
            for word in extras.get("missing_words") or []:
                rows.append({"Defect ID": f"{oid}-MISSING-{len(rows)+1}", "Category": "completeness",
                             "Severity": "HIGH", "Description": "Missing word detected by current lexical alignment",
                             "Expected": word, "Actual": "Missing", "Start s": None, "End s": None,
                             "Evidence": "Review word timeline; exact target-language reference required",
                             "Status": "REVIEW", "Reviewer Note": "Legacy source/dub lexical comparison; verify language pair."})
    return rows

def _audio_rows(results):
    rows = []
    for oid, out in (results.get("outputs") or {}).items():
        aq = (out.get("audio_quality") or (out.get("extras") or {}).get("audio_quality") or {})
        rows.append({"Output ID": oid, "Vendor": out.get("vendor", ""),
                     "Integrated LUFS": aq.get("integrated_lufs", "NOT_MEASURED"),
                     "True Peak dBTP": aq.get("true_peak_dbtp", "NOT_MEASURED"),
                     "Clipped Samples %": aq.get("clipped_samples_pct", "NOT_MEASURED"),
                     "SNR dB": aq.get("snr_db", "NOT_MEASURED"),
                     "Silence Gaps": json.dumps(aq.get("gaps", []), ensure_ascii=False),
                     "Audio Status": _status(aq.get("status"))})
    return rows

def _meta_rows(results):
    return [
        ["Run ID", results.get("run_id")],
        ["Generated", results.get("generated")],
        ["Mode", results.get("mode")],
        ["Thresholds Version", results.get("thresholds_version")],
        ["Source File", (results.get("source") or {}).get("file")],
        ["Source Language", ((results.get("languages") or {}).get("source") or {}).get("resolved")],
        ["Target Language", (results.get("target") or {}).get("language") or results.get("target_lang")],
        ["Git SHA", results.get("git_sha", "NOT_RECORDED")],
        ["Models", json.dumps(results.get("models", {}), ensure_ascii=False)],
        ["Configuration", json.dumps(results.get("configuration", {}), ensure_ascii=False)],
    ]

def _write_sheet(wb, title, headers, rows):
    ws = wb.create_sheet(title)
    if headers:
        ws.append(headers)
        for row in rows:
            ws.append([row.get(h) for h in headers])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="17365D")
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for col in ws.columns:
            vals = [len(str(c.value)) if c.value is not None else 0 for c in list(col)[:150]]
            width = min(max(max(vals, default=10) + 2, 12), 42)
            ws.column_dimensions[get_column_letter(col[0].column)].width = width
        ws.row_dimensions[1].height = 30
        if ws.max_row > 1:
            ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
            tab = Table(displayName="T_" + "".join(ch for ch in title if ch.isalnum())[:20], ref=ref)
            tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True,
                                                showColumnStripes=False)
            ws.add_table(tab)
        # Status cells
        for col_idx, header in enumerate(headers, 1):
            if "Status" in header:
                letter = get_column_letter(col_idx)
                rng = f"{letter}2:{letter}{max(ws.max_row,2)}"
                ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"FAIL"'],
                    fill=PatternFill("solid", fgColor="F4CCCC")))
                ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"REVIEW"'],
                    fill=PatternFill("solid", fgColor="FCE5CD")))
    return ws

def export_run(run_dir: str | Path) -> Path:
    run_dir = Path(run_dir)
    results = _read_json(run_dir / "results.json", None)
    if not results:
        raise FileNotFoundError(f"results.json not found or invalid: {run_dir}")
    report_dir = run_dir / "report"
    report_dir.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Executive Summary"
    summary = results.get("summary") or {}
    outputs = results.get("outputs") or {}
    defects = results.get("defects") or []
    scenarios = results.get("scenario_results") or []
    ws.append(["Field", "Value"])
    summary_rows = [
        ("Run ID", results.get("run_id")),
        ("Generated", results.get("generated")),
        ("Mode", results.get("mode")),
        ("Output Count", len(outputs)),
        ("Scenario Count", len(scenarios)),
        ("Defect Count", len(defects)),
        ("Overall Status", results.get("overall_status", "REVIEW")),
        ("Overall Score", summary.get("overall_score", "NOT_MEASURED")),
        ("Coverage %", summary.get("coverage_pct", "NOT_MEASURED")),
        ("Release Recommendation", results.get("release_recommendation", "REVIEW")),
        ("Status Legend", "PASS / FAIL / REVIEW / NOT_MEASURED / NOT_TESTED"),
        ("Source", (results.get("source") or {}).get("file")),
    ]
    for r in summary_rows: ws.append(list(r))
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="17365D")
    ws.freeze_panes = "A2"; ws.column_dimensions["A"].width = 28; ws.column_dimensions["B"].width = 65
    _write_sheet(wb, "Vendor Metrics",
                 list(_flatten_metrics(results)[0].keys()) if outputs else ["Output ID","Vendor"],
                 _flatten_metrics(results))
    word_rows = _word_rows(run_dir, results)
    _write_sheet(wb, "Word Timeline",
                 list(word_rows[0].keys()) if word_rows else ["Output ID","Word","Word Start s","Word End s","Phoneme Data","Lip-sync Status"],
                 word_rows)
    scenario_rows = _scenario_rows(results)
    _write_sheet(wb, "Scenario Results",
                 list(scenario_rows[0].keys()) if scenario_rows else ["Scenario ID","Category","Status","Expected","Observed","Evidence Start s","Evidence End s"],
                 scenario_rows)
    defect_rows = _defect_rows(results)
    _write_sheet(wb, "Defects",
                 list(defect_rows[0].keys()) if defect_rows else ["Defect ID","Category","Severity","Description","Start s","End s","Status"],
                 defect_rows)
    audio_rows = _audio_rows(results)
    _write_sheet(wb, "Audio Quality",
                 list(audio_rows[0].keys()) if audio_rows else ["Output ID","Integrated LUFS","True Peak dBTP","Clipped Samples %","SNR dB","Audio Status"],
                 audio_rows)
    caption_rows = results.get("caption_results", []) or []
    _write_sheet(wb, "Captions",
                 list(caption_rows[0].keys()) if caption_rows else ["Caption ID","Text","Start s","End s","Status"],
                 caption_rows)
    bench_rows = results.get("competitor_benchmark", []) or []
    _write_sheet(wb, "Competitor Benchmark",
                 list(bench_rows[0].keys()) if bench_rows else ["Platform","Mode","Language Pair","Status","Score","Processing Time s","Cost","Evidence"],
                 bench_rows)
    meta = wb.create_sheet("Run Metadata")
    for row in _meta_rows(results): meta.append(row)
    meta.column_dimensions["A"].width = 26; meta.column_dimensions["B"].width = 85
    for c in meta["A"]:
        c.font = Font(bold=True, color="17365D")
    xlsx = report_dir / "QA_Results.xlsx"
    wb.save(xlsx)
    return xlsx
